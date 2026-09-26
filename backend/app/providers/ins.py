"""INSDataProvider — Institut National de la Statistique (Tunisie).

Source: monthly publication "Commerce extérieur aux prix courants" (official customs-based
foreign-trade statistics). Each publication ships XLSX files with YEAR-TO-DATE cumulative
values for 3 years: by HS chapter (value TND + weight KG), by country (MTND) and totals.
Monthly values are derived as cum(m) - cum(m-1) using the most recent publication for each
(year, month) — a transformation recorded in the data lineage.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

import openpyxl
import pandas as pd
from sqlalchemy.orm import Session

from .. import cache
from ..models import CustomsRecord, EconomicIndicator
from .base import CONNECTED, ERROR, DataProvider, ProviderStatus

BASE = "https://www.ins.tn"
MONTHS_FR = ["janvier", "fevrier", "mars", "avril", "mai", "juin", "juillet", "aout",
             "septembre", "octobre", "novembre", "decembre"]
HDR = re.compile(r"(export|import)\s*(valeur|poids)?\s*(\d{1,2})\s*mois\s*(\d{4})", re.I)
# Regional / residual rows published next to countries (must not be matched to a country).
# Keeps real countries such as PAYS-BAS, AFRIQUE DU SUD, UNION DES COMORES.
AGGREGATE_ROWS = re.compile(
    r"^(pays[ _](?!bas)|total|ensemble|autres?(\s|$)|autre pays|u\.?e(\s|$)|uma(\s|$)|zone|"
    r"afrique(?! du sud)|asie|am[ée]rique|europe|oc[ée]anie|maghreb|reste|avitaillement|non d[ée]termin)",
    re.I)


def _flow(word: str) -> str:
    return "M" if word.lower().startswith("i") else "X"


class INSDataProvider(DataProvider):
    key = "ins_trade"
    name = "INS Tunisie — Commerce extérieur aux prix courants"
    provider = "INSDataProvider"
    url = "https://www.ins.tn/publication"
    access_type = "PUBLIC_OFFICIAL_STATISTICS"
    license = "Institut National de la Statistique — publication officielle"

    def __init__(self, start_year: int = 2023):
        self.start_year = start_year

    # ------------------------------------------------------------------ health
    def health(self) -> ProviderStatus:
        ok = cache.exists(BASE + "/")
        return ProviderStatus(self.key, CONNECTED if ok else ERROR,
                              "ins.tn reachable" if ok else "ins.tn unreachable")

    # ------------------------------------------------------------------ discovery
    def _publication_links(self, year: int, month: int) -> dict[str, str] | None:
        slugs = [f"commerce-exterieur-aux-prix-courants-{MONTHS_FR[month-1]}-{year}",
                 f"commerce-exterieur-au-prix-courant-{MONTHS_FR[month-1]}-{year}"]
        for slug in slugs:
            url = f"{BASE}/publication/{slug}"
            try:
                # past publications do not change: cache 30 days; recent months 1 day
                age = 24 if (datetime.utcnow().year == year) else 24 * 30
                resp = cache.fetch(url, max_age_hours=age, min_interval=0.7)
            except Exception:
                continue
            links = sorted(set(re.findall(r'href="([^"]+\.xlsx)"', resp.text(), re.I)))
            if not links:
                continue
            out = {"page": url}
            for l in links:
                n = l.rsplit("/", 1)[-1].lower()
                full = l if l.startswith("http") else BASE + l
                if "chapit" in n:
                    out["chapter"] = full
                elif "pays" in n:
                    out["country"] = full
                elif "comext" in n:
                    out["global"] = full
            return out
        return None

    # ------------------------------------------------------------------ parsers
    @staticmethod
    def _load(content: bytes):
        return openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)

    @staticmethod
    def _header(rows):
        for i, r in enumerate(rows[:12]):
            cells = [str(c) if c is not None else "" for c in r]
            if sum(1 for c in cells if HDR.search(c)) >= 2:
                return i, cells
        return None, None

    def parse_chapter(self, content: bytes) -> list[dict]:
        ws = self._load(content).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        hi, cells = self._header(rows)
        if hi is None:
            return []
        cols = {}
        for j, c in enumerate(cells):
            m = HDR.search(c)
            if m:
                cols[j] = (_flow(m.group(1)), (m.group(2) or "VALEUR").upper(), int(m.group(3)), int(m.group(4)))
        code_j = next((j for j, c in enumerate(cells) if c.strip().upper() == "CODE"), 0)
        lib_j = next((j for j, c in enumerate(cells) if "LIBEL" in c.upper()), 1)
        out: dict[tuple, dict] = {}
        for r in rows[hi + 1:]:
            code = r[code_j] if code_j < len(r) else None
            if code is None:
                continue
            code = str(code).strip()
            if not re.fullmatch(r"\d{1,2}", code):
                continue
            code = code.zfill(2)
            lib = str(r[lib_j]).strip() if r[lib_j] else None
            for j, (flow, measure, months, year) in cols.items():
                v = r[j] if j < len(r) else None
                if v is None or isinstance(v, str):
                    continue
                k = (flow, code, year, months)
                d = out.setdefault(k, {"flow": flow, "hs": code, "desc": lib, "year": year, "months": months})
                d["value_tnd" if measure == "VALEUR" else "weight_kg"] = float(v)
        return list(out.values())

    def parse_country(self, content: bytes) -> list[dict]:
        ws = self._load(content).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        hi, cells = self._header(rows)
        if hi is None:
            return []
        name_j = next((j for j, c in enumerate(cells) if c.strip().lower() in ("libpay", "pays")), 1)
        cols = {j: HDR.search(c) for j, c in enumerate(cells) if HDR.search(c)}
        out = []
        for r in rows[hi + 1:]:
            name = r[name_j] if name_j < len(r) else None
            if not name or not isinstance(name, str):
                continue
            name = name.strip()
            for j, m in cols.items():
                v = r[j] if j < len(r) else None
                if v is None or isinstance(v, str):
                    continue
                out.append({"flow": _flow(m.group(1)), "country": name,
                            "aggregate": bool(AGGREGATE_ROWS.search(name)),
                            "year": int(m.group(4)), "months": int(m.group(3)), "value_mtnd": float(v)})
        return out

    def parse_global(self, content: bytes) -> list[dict]:
        wb = self._load(content)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        years, months = None, None
        out, seen = [], set()
        for r in rows:
            cells = [c for c in r]
            txt = [str(c) for c in cells if c is not None]
            found = [re.search(r"(\d{1,2})\s*mois\s*(\d{4})", t, re.I) for t in txt]
            found = [f for f in found if f]
            if len(found) >= 2 and years is None:
                years = [(int(f.group(1)), int(f.group(2))) for f in found]
                continue
            if years and cells and any(isinstance(c, str) and c.strip().lower() in ("exportations", "importations") for c in cells):
                label = next(c for c in cells if isinstance(c, str)).strip().lower()
                if label in seen:  # only the first block = national total ("ENSEMBLE"); later blocks are regimes
                    continue
                seen.add(label)
                nums = [c for c in cells if isinstance(c, (int, float))][: len(years)]
                for (mo, yr), v in zip(years, nums):
                    out.append({"flow": "X" if label.startswith("export") else "M", "year": yr, "months": mo, "value_mtnd": float(v)})
        return out

    # ------------------------------------------------------------------ ingest
    def ingest(self, db: Session, log) -> ProviderStatus:
        now = datetime.utcnow()
        chapter_rows, country_rows, global_rows = [], [], []
        pubs = 0
        latest = None
        for year in range(self.start_year, now.year + 1):
            for month in range(1, 13):
                if year == now.year and month > now.month:
                    break
                links = self._publication_links(year, month)
                if not links:
                    continue
                pubs += 1
                pub_rank = year * 100 + month  # newer publication wins for revisions
                latest = max(latest or 0, pub_rank)
                for kind, parser, bucket in (("chapter", self.parse_chapter, chapter_rows),
                                             ("country", self.parse_country, country_rows)):
                    if kind not in links:
                        continue
                    try:
                        resp = cache.fetch(links[kind], max_age_hours=24 * 30, min_interval=0.7)
                        self.register_raw(db, resp, period=f"{year}-{month:02d} YTD")
                        for row in parser(resp.content):
                            row.update(pub=pub_rank, url=links[kind], retrieved_at=resp.retrieved_at)
                            bucket.append(row)
                    except Exception as e:  # keep going; the error is logged, never masked with fake data
                        log(f"INS {kind} {year}-{month}: {e}")
        if not chapter_rows:
            st = ProviderStatus(self.key, ERROR, "No INS publication could be parsed")
            self.save_status(db, st)
            return st

        db.query(CustomsRecord).filter(CustomsRecord.source_key == self.key).delete()
        db.query(EconomicIndicator).filter(EconomicIndicator.source_key == self.key).delete()

        n = 0
        # ---- chapters: cumulative -> monthly
        ch = pd.DataFrame(chapter_rows)
        ch = ch.sort_values("pub").groupby(["flow", "hs", "year", "months"], as_index=False).last()
        for (flow, hs, year), g in ch.groupby(["flow", "hs", "year"]):
            g = g.set_index("months").sort_index()
            for m in g.index:
                prev = 0.0 if m == 1 else (g.loc[m - 1, "value_tnd"] if (m - 1) in g.index else None)
                prevw = 0.0 if m == 1 else (g.loc[m - 1, "weight_kg"] if (m - 1) in g.index and "weight_kg" in g else None)
                if prev is None:
                    continue
                row = g.loc[m]
                val = row["value_tnd"] - prev
                w = (row.get("weight_kg") - prevw) if (prevw is not None and pd.notna(row.get("weight_kg"))) else None
                db.add(CustomsRecord(dataset="ins_chapter_month", flow=flow, period=f"{year}-{m:02d}", period_type="M",
                                     hs_code=hs, hs_description=row["desc"], value=float(val), value_unit="TND",
                                     weight_kg=float(w) if w is not None else None, source_key=self.key,
                                     source_url=row["url"], retrieved_at=row["retrieved_at"]))
                n += 1
        # ---- countries: cumulative -> monthly (MTND)
        co = pd.DataFrame(country_rows)
        co = co.sort_values("pub").groupby(["flow", "country", "year", "months"], as_index=False).last()
        for (flow, country, year), g in co.groupby(["flow", "country", "year"]):
            g = g.set_index("months").sort_index()
            for m in g.index:
                prev = 0.0 if m == 1 else (g.loc[m - 1, "value_mtnd"] if (m - 1) in g.index else None)
                if prev is None:
                    continue
                row = g.loc[m]
                db.add(CustomsRecord(dataset="ins_country_month_agg" if row["aggregate"] else "ins_country_month",
                                     flow=flow, period=f"{year}-{m:02d}", period_type="M",
                                     partner_name=country, value=float(row["value_mtnd"] - prev), value_unit="MTND",
                                     source_key=self.key, source_url=row["url"], retrieved_at=row["retrieved_at"]))
                n += 1
        # ---- totals: monthly and annual indicators
        # National totals = sum of the 97 HS chapters of the same publication (the chapter file covers all
        # goods). Checked against the published "Comext" totals (see tests). Documented in the lineage.
        gl = ch.groupby(["flow", "year", "months"], as_index=False).agg(
            value_tnd=("value_tnd", "sum"), pub=("pub", "max"), url=("url", "first"), retrieved_at=("retrieved_at", "first"))
        gl["value_mtnd"] = gl["value_tnd"] / 1e6
        label = {"M": ("ins_imports", "Imports of goods (INS)"), "X": ("ins_exports", "Exports of goods (INS)")}
        for (flow, year), g in gl.groupby(["flow", "year"]):
            g = g.set_index("months").sort_index()
            ind, lab = label[flow]
            for m in g.index:
                prev = 0.0 if m == 1 else (g.loc[m - 1, "value_mtnd"] if (m - 1) in g.index else None)
                if prev is None:
                    continue
                db.add(EconomicIndicator(indicator=ind + "_monthly", label=lab + " — monthly", period=f"{year}-{m:02d}",
                                         frequency="M", value=float(g.loc[m, "value_mtnd"] - prev), unit="MTND",
                                         source_key=self.key, source_url=g.loc[m, "url"], retrieved_at=g.loc[m, "retrieved_at"]))
            mx = int(g.index.max())
            db.add(EconomicIndicator(indicator=ind + ("_annual" if mx == 12 else "_ytd"), label=lab + (" — annual" if mx == 12 else f" — Jan–{mx:02d} YTD"),
                                     period=str(year) if mx == 12 else f"{year}-YTD{mx:02d}", frequency="A",
                                     value=float(g.loc[mx, "value_mtnd"]), unit="MTND", source_key=self.key,
                                     source_url=g.loc[mx, "url"], retrieved_at=g.loc[mx, "retrieved_at"]))
        db.flush()
        per = ch.assign(p=ch.year * 100 + ch.months)
        lo, hi = per.p.min(), per.p.max()
        st = ProviderStatus(self.key, CONNECTED, f"{pubs} monthly publications parsed", n,
                            period=f"{str(lo)[:4]}-{str(lo)[4:]} → {str(hi)[:4]}-{str(hi)[4:]}",
                            last_updated=now)
        self.save_status(db, st)
        return st
